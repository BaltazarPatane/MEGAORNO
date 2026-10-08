"""Integrity checks at file, acquisition and export boundaries."""
import csv
import io
import math
import os
import queue
import shutil
import subprocess
import sys
import tempfile
import threading
import unittest
from datetime import datetime, timedelta, time
from pathlib import Path
from unittest.mock import Mock, patch
from types import SimpleNamespace

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from openpyxl import Workbook, load_workbook
from matplotlib.figure import Figure
from charts import render_chart
from live import WcfBridge, decodificar_madgetech, redact_diagnostic
from models import Registro, Sesion, cargar_xlsx, exportar_csv
from reports import exportar_excel, exportar_pdf


class FileIntegrity(unittest.TestCase):
    def make_excel(self,path,rows):
        book=Workbook();sheet=book.active
        sheet.append([None,None,'Canal 2'])
        sheet.append(['Fecha','Hora','Termopar 1 (°C)'])
        for row in rows:sheet.append(row)
        book.save(path);book.close()

    def test_textual_date_keeps_time_column_and_real_intervals(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'temperaturas.xlsx'
            self.make_excel(path,[['19/12/2024','08:00:00',20],['2024-12-19','08:00:15',25],
                                  [datetime(2024,12,19,8,0,31),time(12,0),29]])
            record=cargar_xlsx(path);channel=record.canales[0]
            self.assertEqual([t for t,_ in record.muestras],[datetime(2024,12,19,8,0),datetime(2024,12,19,8,0,15),datetime(2024,12,19,8,0,31)])
            self.assertEqual(record.gradientes[1][channel]*60,1200)
            self.assertEqual(record.gradientes[2][channel]*60,900)

    def test_malformed_time_is_rejected_instead_of_midnight(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'bad.xlsx'
            self.make_excel(path,[[datetime(2024,12,19),123,20]])
            with self.assertRaisesRegex(ValueError,'Fila 3'):
                cargar_xlsx(path)

    def test_nan_inf_and_error_cells_are_missing_with_notes(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'bad_values.xlsx'
            self.make_excel(path,[[datetime(2024,12,19),time(8,0),'NaN'],[datetime(2024,12,19),time(8,0,15),'inf'],[datetime(2024,12,19),time(8,0,30),'#VALUE!']])
            record=cargar_xlsx(path)
            self.assertEqual(len(record.notas),3)
            self.assertTrue(all(list(values.values())==[None] for _,values in record.muestras))

    def test_failed_atomic_replace_keeps_previous_session_and_cleans_temp(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'captura.horno.json';session=Sesion(titulo='Anterior');session.guardar(path)
            before=path.read_bytes();session.titulo='Nuevo'
            with patch('models.os.replace',side_effect=OSError('Disco no disponible')):
                with self.assertRaises(OSError):session.guardar(path)
            self.assertEqual(path.read_bytes(),before)
            self.assertEqual(list(Path(folder).glob('*.tmp')),[])
            self.assertEqual(Sesion.abrir(path).titulo,'Anterior')

    def test_nonfinite_save_keeps_previous_capture(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'captura.horno.json';session=Sesion();session.guardar(path)
            before=path.read_bytes();session.registro=Registro(canales=['TC1'],muestras=[(datetime(2026,1,1),{'TC1':math.nan})])
            with self.assertRaises(ValueError):session.guardar(path)
            self.assertEqual(path.read_bytes(),before)
            self.assertEqual(list(Path(folder).glob('*.tmp')),[])


class TransportIntegrity(unittest.TestCase):
    def payload(self):
        return dict(serial='R59022',timestamp='2026-10-06T09:00:00',channels=[dict(position=2,unit='Fahrenheit',status='OK',value='1166')])

    def test_conversion_does_not_overflow_an_intermediate_product(self):
        payload=self.payload();payload['channels'][0]['value']='1e308'
        rows,errors=decodificar_madgetech(payload)
        self.assertFalse(errors);self.assertTrue(math.isfinite(rows[0][2]))
        self.assertAlmostEqual(rows[0][2]/1e308,5/9)

    def test_malformed_payloads_fail_explicitly(self):
        for value in [None,[],{},dict(self.payload(),serial=None),dict(self.payload(),channels=[None])]:
            with self.subTest(value=value),self.assertRaises(ValueError):decodificar_madgetech(value)

    def test_authentication_diagnostics_hide_credentials(self):
        config=dict(username='operador_123',password='Secreto$321')
        diagnostic={'message':'Authenticate operador_123 failed Secreto$321','nested':['Secreto$321',7]}
        cleaned=redact_diagnostic(diagnostic,config)
        self.assertNotIn(config['username'],str(cleaned));self.assertNotIn(config['password'],str(cleaned))
        self.assertEqual(cleaned['nested'][1],7)
        self.assertIn(config['password'],str(diagnostic))

    def test_stop_process_exit_race_does_not_escape(self):
        bridge=WcfBridge(queue.Queue());bridge.process=Mock()
        bridge.process.poll.return_value=None;bridge.process.terminate.side_effect=ProcessLookupError('Already exited')
        bridge.stop()
        self.assertTrue(bridge.stop_event.is_set());self.assertEqual(bridge.request_id,1)

    def test_cancelled_and_new_bridges_do_not_compile_the_same_file_concurrently(self):
        entered=threading.Event();release=threading.Event()
        with tempfile.TemporaryDirectory() as folder:
            target=Path(folder);compiler=target/'Windows/Microsoft.NET/Framework/v4.0.30319/csc.exe'
            compiler.parent.mkdir(parents=True);compiler.touch()
            old=WcfBridge(queue.Queue());new=WcfBridge(queue.Queue())
            def compile_command(command,**kwargs):
                entered.set();release.wait(2)
                Path(next(item[5:] for item in command if item.startswith('/out:'))).write_bytes(b'test executable')
                return SimpleNamespace(returncode=0,stdout='',stderr='')
            def fake_process(*args,**kwargs):
                process=Mock();process.stdin=io.StringIO();process.stdout=io.StringIO('')
                process.poll.return_value=0;process.wait.return_value=0
                return process
            runner=Mock(side_effect=compile_command)
            platform=SimpleNamespace(name='nt',environ={'WINDIR':str(target/'Windows')})
            processes=SimpleNamespace(run=runner,Popen=Mock(side_effect=fake_process),CREATE_NO_WINDOW=0,
                                      PIPE=subprocess.PIPE,STDOUT=subprocess.STDOUT,TimeoutExpired=subprocess.TimeoutExpired)
            with patch('live.os',platform),patch('live.subprocess',processes):
                old.start({},target/'build')
                try:
                    self.assertTrue(entered.wait(2));old.stop();new.start({},target/'build')
                finally:release.set();old.thread.join(3)
                if new.thread:new.thread.join(3)
            self.assertEqual(runner.call_count,1)
            self.assertFalse(old.running);self.assertFalse(new.running)
            events=[]
            while not new.events.empty():events.append(new.events.get())
            self.assertFalse([event for event in events if event['type']=='error'])
            self.assertEqual(events[-1]['type'],'finished')

    def test_duplicate_start_is_rejected_and_stop_invalidates_generation(self):
        bridge=WcfBridge(queue.Queue());finish=threading.Event()
        with patch.object(bridge,'_run',side_effect=lambda *args:finish.wait(2)):
            bridge.start({},ROOT)
            try:
                generation=bridge.request_id
                with self.assertRaises(ValueError):bridge.start({},ROOT)
                bridge.stop();self.assertGreater(bridge.request_id,generation)
            finally:finish.set();bridge.thread.join(2)
        self.assertFalse(bridge.running)


class ExportIntegrity(unittest.TestCase):
    def session(self):
        session=Sesion(registro=cargar_xlsx(ROOT/'ejemplos'/'R59022 MultiChannel.xlsx'),titulo='Ensayo de integridad',modo='Ejemplo · SIMULACIÓN')
        session.perfil.desfase=17
        return session

    def test_excel_alias_header_is_text_and_values_remain_exact(self):
        session=self.session();session.alias[session.registro.canales[0]]='=HYPERLINK("bad")'
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'ensayo.xlsx';exportar_excel(session,path)
            book=load_workbook(path,data_only=False)
            try:
                self.assertEqual(book['Datos']['E1'].data_type,'s')
                self.assertEqual(book['Datos']['E1'].value,'=HYPERLINK("bad") °C')
                self.assertEqual(book['Datos']['F3'].value,session.registro.gradientes[1][session.registro.canales[0]]*60)
                self.assertFalse(book._external_links)
                self.assertTrue(all(cell.data_type!='f' for sheet in book for row in sheet for cell in row))
                self.assertEqual(book['Estadísticas']['E2'].value,652.5)
                self.assertEqual(book['Estadísticas']['D2'].value,session.registro.muestras[-1][0])
            finally:book.close()

    def test_excel_plot_breaks_long_gaps_and_preserves_asynchronous_channels(self):
        t=datetime(2026,10,6,9);session=Sesion(registro=Registro())
        session.registro.agregar([(t,'TC1',20),(t+timedelta(seconds=15),'TC2',30),
                                  (t+timedelta(minutes=5),'TC1',25),(t+timedelta(minutes=20),'TC1',30),
                                  (t+timedelta(minutes=21),'TC1',None),(t+timedelta(minutes=22),'TC1',35)])
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'gaps.xlsx';exportar_excel(session,path);book=load_workbook(path)
            try:
                self.assertEqual(book['Datos'].max_row,7)
                drawing=book['Trazado'];self.assertEqual(drawing.sheet_state,'hidden')
                self.assertEqual([(drawing.cell(i,3).value,drawing.cell(i,4).value) for i in range(2,8)],
                                 [(0,20),(5,25),(20,None),(20,30),(21,None),(22,35)])
                chart=book['Curvas']._charts[0]
                self.assertEqual(chart.display_blanks,'gap');self.assertFalse(chart.visible_cells_only)
                self.assertIn('Trazado',chart.series[1].xVal.numRef.f)
            finally:book.close()

    def test_csv_uses_measurement_times_and_missing_rates_stay_blank(self):
        t=datetime(2026,10,6,9)
        session=Sesion(registro=Registro())
        session.registro.agregar([(t,'TC1',20),(t+timedelta(seconds=15),'TC1',25),
                                  (t+timedelta(seconds=30),'TC1',None),(t+timedelta(seconds=45),'TC1',27),
                                  (t+timedelta(minutes=20),'TC1',30)])
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'ensayo.csv';exportar_csv(session,path)
            self.assertTrue(path.read_bytes().startswith(b'\xef\xbb\xbf'))
            with path.open(encoding='utf-8-sig',newline='') as stream:rows=list(csv.reader(stream,delimiter=';'))
            self.assertEqual(rows[2][0],'2026-10-06 09:00:15')
            self.assertEqual(float(rows[2][5]),1200);self.assertEqual(float(rows[2][6]),20)
            self.assertEqual(rows[3][4:7],['','',''])
            self.assertEqual(rows[4][5:7],['','']);self.assertEqual(rows[5][5:7],['',''])

    def test_pdf_contains_complete_recipe_latest_dates_and_all_channels(self):
        session=self.session();figure=Figure(figsize=(10,5));a,b=figure.subplots(2,1,sharex=True)
        render_chart(figure,a,b,session,list(range(len(session.registro.canales))),mode='Ambas',show_ideal=True,show_band=True,full_program=True)
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'ensayo.pdf';exportar_pdf(session,figure,path)
            self.assertTrue(path.read_bytes().startswith(b'%PDF-'))
            self.assertGreater(path.stat().st_size,15000)
            if shutil.which('pdftotext'):
                result=subprocess.run(['pdftotext','-layout',str(path),'-'],check=True,capture_output=True,text=True)
                text=result.stdout
                for expected in ['Ensayo de integridad','SIMULACIÓN','17 min','Termocuplas · 1','Termocuplas · 3','19/12/2024 15:29:11','652.50','Última válida']:
                    self.assertIn(expected,text)
        self.assertEqual(len(figure.axes),2)


if __name__=='__main__':unittest.main()
