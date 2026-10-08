"""Pause/race regressions, persistent gaps and actual Qt mouse-wheel gestures."""
from copy import deepcopy
from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch
import csv
import io
import json
import os
import queue
import sys
import tempfile
import unittest

os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from PySide6.QtCore import QPoint, QPointF, Qt
from PySide6.QtGui import QWheelEvent
from PySide6.QtWidgets import QApplication
from matplotlib.figure import Figure
from openpyxl import load_workbook
from acquisition import AcquisitionController, ConnectionSettings
from charts import render_chart
from gui import App
from live import WcfBridge
from models import Registro, Sesion, exportar_csv
from reports import exportar_excel
from ui.chart_widget import ChartWidget
from test_acquisition import FakeBridge, reading
from test_gui_workflow import WorkflowController, payload


class PauseAcquisition(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.path=Path(self.temp.name)/'capture.horno.json'
        self.controller=AcquisitionController(self.temp.name,FakeBridge)
        self.addCleanup(self.controller.disconnect)
        self.settings=ConnectionSettings(vendor_dir=self.temp.name)
        self.controller.connect(self.settings,self.path)
        self.bridge=self.controller.bridge
        self.bridge.emit('mt_devices',['R59022'])
        for minute in (0,1):self.bridge.emit('mt_reading',reading(f'2026-10-06T08:0{minute}:00',str(630+minute)))
        self.controller.drain()

    def test_pause_preserves_buffer_file_raw_response_and_session_even_with_late_frames(self):
        c=self.controller;session=c.session;record=session.registro
        before=deepcopy(record);disk=self.path.read_bytes();raw=Path(str(self.path)+'.wcf.jsonl').read_bytes()
        c.pause();epoch=self.bridge.epoch;c.pause()
        for _ in range(30):
            self.bridge.emit('mt_reading',{'malformed':'must never be parsed'})
            self.bridge.emit('mt_waiting','Late waiting status')
        notifications=c.drain()
        self.assertTrue(c.paused);self.assertTrue(c.running);self.assertTrue(self.bridge.running)
        self.assertFalse(self.bridge.stopped);self.assertEqual(self.bridge.starts,1)
        self.assertEqual(self.bridge.epoch,epoch);self.assertIs(c.session,session)
        self.assertIs(c.session.registro,record);self.assertEqual(record,before)
        self.assertEqual(self.path.read_bytes(),disk)
        self.assertEqual(Path(str(self.path)+'.wcf.jsonl').read_bytes(),raw)
        self.assertFalse([e for e in notifications if e['type'] in ('sample','ready','error')])
        with self.assertRaises(ValueError):c.connect(self.settings,self.path)
        with self.assertRaises(ValueError):c.search(self.settings)

    def test_rapid_resume_rejects_old_epochs_and_times_then_appends_with_persistent_gap(self):
        c=self.controller;session=c.session;before=deepcopy(session.registro.muestras)
        gradients=deepcopy(session.registro.gradientes)
        c.pause();c.resume();c.pause();c.resume()
        for epoch in range(self.bridge.epoch):
            self.bridge.emit('mt_reading',reading('2026-10-06T08:02:00','999'),epoch=epoch)
            self.bridge.emit('mt_waiting','stale',epoch=epoch)
        self.bridge.emit('mt_reading',reading('2026-10-06T08:01:00','999'))
        self.bridge.emit('mt_reading',reading('2026-10-06T08:03:00','632'))
        self.bridge.emit('mt_reading',reading('2026-10-06T08:04:00','634'))
        events=c.drain();r=session.registro;channel=r.canales[0]
        self.assertIs(c.session,session);self.assertEqual(r.muestras[:2],before)
        self.assertEqual(r.gradientes[:2],gradients);self.assertEqual(len(r.muestras),4)
        self.assertIsNone(r.gradientes[2][channel]);self.assertEqual(r.gradientes[3][channel],2)
        self.assertEqual(r.cortes,[datetime(2026,10,6,8,3)])
        self.assertEqual(c.state,'receiving');self.assertEqual(self.bridge.starts,1)
        self.assertFalse([e for e in events if e['type'] in ('ready','error')])
        self.assertEqual(Sesion.abrir(self.path).registro,r)
        self.assertEqual(json.loads(self.path.read_text())['version'],3)
        self.assertEqual(len(Path(str(self.path)+'.wcf.jsonl').read_text().splitlines()),4)

    def test_failed_control_and_transport_exit_preserve_capture_and_stop_explicitly(self):
        before=deepcopy(self.controller.session.registro)
        with patch.object(self.bridge,'set_paused',side_effect=BrokenPipeError('Bridge exited')):
            self.controller.pause()
        events=self.controller.drain()
        self.assertEqual(self.controller.state,'error');self.assertFalse(self.controller.running)
        self.assertEqual(self.controller.session.registro,before)
        self.assertTrue(any(e['type']=='error' for e in events))
        self.assertEqual(Sesion.abrir(self.path).registro,before)

    def test_resume_failure_does_not_replace_history_or_claim_active_reception(self):
        self.controller.pause();before=self.path.read_bytes()
        with patch.object(self.bridge,'set_paused',side_effect=OSError('Control unavailable')):
            self.controller.resume()
        self.assertEqual(self.controller.state,'error');self.assertEqual(self.path.read_bytes(),before)
        self.assertEqual(len(self.controller.session.registro.muestras),2)

    def test_save_failure_after_resume_keeps_new_sample_and_gap_in_memory(self):
        c=self.controller;c.pause();c.resume();before=self.path.read_bytes()
        self.bridge.emit('mt_reading',reading('2026-10-06T08:02:00','633'))
        with patch.object(c.session,'guardar',side_effect=OSError('Disk full')):c.drain()
        self.assertEqual(c.state,'error');self.assertEqual(self.path.read_bytes(),before)
        self.assertEqual(len(c.session.registro.muestras),3)
        self.assertEqual(c.session.registro.cortes,[datetime(2026,10,6,8,2)])
        copy=Path(self.temp.name)/'recovery.horno.json';c.session.guardar(copy)
        self.assertEqual(Sesion.abrir(copy).registro,c.session.registro)


class GapIntegrity(unittest.TestCase):
    def test_pause_breaks_async_channels_band_streak_and_all_export_plots(self):
        t=datetime(2026,10,6,8);r=Registro()
        rows=[(t,'A',630),(t,'B',630),(t+timedelta(minutes=1),'A',631),
              (t+timedelta(minutes=1),'B',631),(t+timedelta(minutes=2),'A',632),
              (t+timedelta(minutes=3),'B',633),(t+timedelta(minutes=4),'A',634),
              (t+timedelta(minutes=4),'B',634)]
        r.cortes=[t+timedelta(minutes=2)];r.agregar(rows)
        session=Sesion(registro=r);summary,_=session.analisis()
        self.assertIsNone(r.gradientes[2]['A']);self.assertIsNone(r.gradientes[3]['B'])
        self.assertEqual(r.gradientes[4]['A'],1);self.assertEqual(r.gradientes[4]['B'],1)
        self.assertEqual([s['racha'] for s in summary],[2,1])
        fig=Figure();ax_t,ax_g=fig.subplots(2,1,sharex=True)
        render_chart(fig,ax_t,ax_g,session,[0,1],mode='Ambas')
        self.assertIsNone(ax_t.lines[0].get_ydata()[2])
        self.assertIsNone(ax_t.lines[1].get_ydata()[2])
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder);exportar_excel(session,p/'gap.xlsx');exportar_csv(session,p/'gap.csv')
            book=load_workbook(p/'gap.xlsx')
            try:
                self.assertIn('Gradiente',book['Datos']['F1'].value)
                self.assertIsNone(book['Datos']['F4'].value)
                self.assertIsNone(book['Datos']['I5'].value)
                self.assertIsNone(book['Trazado']['D4'].value)
                self.assertIsNone(book['Trazado']['F4'].value)
                self.assertEqual(book['Datos'].max_row,6)
            finally:book.close()
            with (p/'gap.csv').open(encoding='utf-8-sig') as stream:table=list(csv.reader(stream,delimiter=';'))
            self.assertIn('Gradiente',table[0][5]);self.assertEqual(table[3][5],'')

    def test_original_v2_without_gap_field_still_opens_with_identical_gradients(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'old.horno.json';session=Sesion()
            session.registro.agregar([(datetime(2026,1,1,8),'TC1',20),(datetime(2026,1,1,8,1),'TC1',22)])
            session.guardar(path);data=json.loads(path.read_text());data['registro'].pop('cortes')
            path.write_text(json.dumps(data))
            restored=Sesion.abrir(path)
            self.assertEqual(restored.registro.gradientes,session.registro.gradientes)
            self.assertEqual(restored.registro.cortes,[])

    def test_control_commands_use_open_stdin_and_leave_process_alive(self):
        bridge=WcfBridge(queue.Queue());bridge.process=Mock()
        bridge.process.poll.return_value=None;bridge.process.stdin=io.StringIO()
        bridge.set_paused(True,1);bridge.set_paused(False,2)
        commands=[json.loads(line) for line in bridge.process.stdin.getvalue().splitlines()]
        self.assertEqual(commands,[{'command':'pause','epoch':1},{'command':'resume','epoch':2}])
        self.assertFalse(bridge.process.stdin.closed);bridge.process.terminate.assert_not_called()


class UiRevision32(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.qt=QApplication.instance() or QApplication([])
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.controller=WorkflowController(self.temp.name)
        self.w=App(self.temp.name,self.controller);self.w.timer.stop();self.w.refresh_timer.stop()
        self.w.login.vendor_dir.setText(self.temp.name)
    def tearDown(self):
        self.w.close();self.w.deleteLater();self.qt.processEvents();self.temp.cleanup()
    def receive(self):
        self.w.connect_device();b=self.controller.bridge;b.emit('mt_devices',['R59022'])
        b.emit('mt_reading',payload());self.w.consume();return b

    def test_pause_button_keeps_tab_selection_view_and_export_and_save_target(self):
        b=self.receive();self.w.flush_refresh();session=self.w.session;path=self.w.session_path
        chart=self.w.curves.chart;chart.set_follow(False)
        limits=(chart.ax_t.get_xlim(),chart.ax_t.get_ylim());self.w.tabs.setCurrentWidget(self.w.data)
        self.w.pause_button.click()
        self.assertTrue(self.controller.paused);self.assertEqual(self.w.pause_button.text(),'Reanudar')
        self.assertIn('PAUSADA',self.w.connection_badge.text());self.assertTrue(self.w.export_button.isEnabled())
        self.assertIs(self.w.tabs.currentWidget(),self.w.data)
        b.emit('mt_reading',payload('2026-09-29T08:01:00','999'));self.w.consume()
        self.assertEqual(len(session.registro.muestras),1)
        copy=Path(self.temp.name)/'copy.horno.json'
        with patch('gui.QFileDialog.getSaveFileName',return_value=(str(copy),'')):self.w.save_copy()
        self.assertEqual(self.w.session_path,path);self.assertEqual(self.controller.session_path,path)
        self.w.pause_button.click();b.emit('mt_reading',payload('2026-09-29T08:02:00','632'))
        self.w.consume();self.w.tabs.setCurrentWidget(self.w.curves);self.w.flush_refresh()
        self.assertIs(self.w.session,session);self.assertEqual(len(session.registro.muestras),2)
        self.assertEqual((chart.ax_t.get_xlim(),chart.ax_t.get_ylim()),limits)
        self.assertEqual(len(Sesion.abrir(copy).registro.muestras),1)
        self.assertEqual(self.w.pause_button.text(),'Pausa')

    def test_logout_during_pause_saves_and_late_events_do_not_reopen_session(self):
        b=self.receive();path=self.w.session_path;self.w.toggle_pause();self.w.logout()
        b.emit('mt_reading',payload('2026-09-29T08:02:00','999'));self.w.consume()
        self.assertFalse(self.w.active);self.assertFalse(self.w.pause_button.isEnabled())
        self.assertEqual(len(Sesion.abrir(path).registro.muestras),1)
        self.assertFalse(self.controller.running)

    def test_example_pause_resumes_same_buffer_instead_of_restarting_replay(self):
        self.w.open_example();self.assertFalse(self.w.pause_button.isEnabled())
        self.w.toggle_replay();session=self.w.session;path=self.w.session_path
        self.w.toggle_pause();before=deepcopy(session.registro.muestras);remaining=list(self.w.replay_rows)
        self.w.replay_step();self.assertEqual(session.registro.muestras,before)
        self.assertEqual(self.w.replay_rows,remaining);self.assertTrue(self.w.replay_paused)
        self.w.toggle_replay();self.w.replay_timer.stop();self.w.replay_step()
        self.assertIs(self.w.session,session);self.assertEqual(self.w.session_path,path)
        self.assertEqual(session.registro.muestras[:len(before)],before)
        self.assertEqual(len(session.registro.muestras),len(before)+1)

    def test_actual_qt_wheel_freezes_other_axis_in_all_modes_both_directions(self):
        self.w.open_example();self.w.show();self.qt.processEvents();chart=self.w.curves.chart
        for mode in ('Temperatura','Gradiente','Ambas'):
            chart.set_options(mode=mode);chart.canvas.draw()
            for ax in (chart.ax_t,chart.ax_g):
                if not ax.get_visible():continue
                for ctrl in (False,True):
                    for delta in (-120,120):
                        before_x=ax.get_xlim();before_y=[a.get_ylim() for a in (chart.ax_t,chart.ax_g)]
                        x,y=ax.transData.transform((sum(before_x)/2,sum(ax.get_ylim())/2))
                        local=QPointF(x/chart.canvas.device_pixel_ratio,(chart.canvas.figure.bbox.height-y)/chart.canvas.device_pixel_ratio)
                        event=QWheelEvent(local,local,QPoint(),QPoint(0,delta),Qt.MouseButton.NoButton,
                            Qt.KeyboardModifier.ControlModifier if ctrl else Qt.KeyboardModifier.NoModifier,
                            Qt.ScrollPhase.NoScrollPhase,False)
                        QApplication.sendEvent(chart.canvas,event);self.qt.processEvents()
                        if ctrl:
                            self.assertEqual(chart.ax_t.get_xlim(),before_x);self.assertEqual(chart.ax_g.get_xlim(),before_x)
                            self.assertNotEqual(ax.get_ylim(),before_y[0 if ax is chart.ax_t else 1])
                        else:
                            self.assertNotEqual(ax.get_xlim(),before_x)
                            self.assertEqual([a.get_ylim() for a in (chart.ax_t,chart.ax_g)],before_y)

    def test_hover_and_click_show_larger_exact_gradient_annotation(self):
        self.w.open_example();chart=self.w.curves.chart;chart.set_options(mode='Gradiente');chart.canvas.draw()
        series=chart._series[chart.ax_g][0];x,y=series['x'][0],series['y'][0]
        px,py=chart.ax_g.transData.transform((x,y))
        e=SimpleNamespace(inaxes=chart.ax_g,xdata=x,ydata=y,x=px,y=py,key=None,button=1,dblclick=False)
        chart._on_motion(e);annotation=chart._overlays[chart.ax_g][2]
        self.assertGreaterEqual(annotation.get_fontsize(),12)
        self.assertIn(f'{y:.10g}',annotation.get_text());self.assertIn('Gradiente:',annotation.get_text())
        chart._hide_cursor();chart._on_press(e);chart._on_release(e)
        self.assertTrue(annotation.get_visible());self.assertIn('Gradiente:',annotation.get_text())

    def test_large_tooltip_restores_entire_canvas_without_ghosts_in_short_dual_plot(self):
        self.w.open_example();chart=ChartWidget();self.addCleanup(chart.close)
        chart.set_session(self.w.session);chart.resize(1000,260);chart.show()
        chart.set_options(mode='Ambas');self.qt.processEvents();chart.canvas.draw();self.qt.processEvents()
        # Freeze the layout for a pixel-exact comparison of cursor restoration,
        # independent of constrained_layout's iterative tick positioning.
        chart.figure.set_layout_engine(None);chart.canvas.draw()
        series=chart._series[chart.ax_t][0];x,y=series['x'][40],series['y'][40]
        px,py=chart.ax_t.transData.transform((x,y))
        e=SimpleNamespace(inaxes=chart.ax_t,xdata=x,ydata=y,x=px,y=py)
        before=bytes(chart.canvas.buffer_rgba());chart._on_motion(e)
        self.assertNotEqual(bytes(chart.canvas.buffer_rgba()),before)
        self.assertFalse(chart._overlays[chart.ax_t][2].get_in_layout())
        chart._hide_cursor()
        self.assertEqual(bytes(chart.canvas.buffer_rgba()),before)


if __name__=='__main__':unittest.main()
