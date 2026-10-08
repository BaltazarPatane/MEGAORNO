import sys
import tempfile
import unittest
from datetime import datetime,timedelta
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from models import Registro,Perfil,Sesion,cargar_xlsx,exportar_csv
from live import decodificar
from reports import exportar_excel
from openpyxl import load_workbook


class Calculos(unittest.TestCase):
    def test_archivo_real(self):
        r=cargar_xlsx(ROOT/'ejemplos'/'R59022 MultiChannel.xlsx')
        self.assertEqual(len(r.muestras),93);self.assertEqual(len(r.canales),2);self.assertIn('Termopar 3',r.canales[1])
        resumen,_=Sesion(registro=r).analisis();self.assertEqual([x['maximo'] for x in resumen],[652.5,644.4])
        self.assertEqual(resumen[0]['tmax'],datetime(2024,12,19,12,44,11))
        self.assertAlmostEqual(r.velocidades[1][r.canales[0]]*60,14.1*3600/301)
        self.assertAlmostEqual(resumen[0]['subida'],6.518272425249,places=7)
        self.assertAlmostEqual(resumen[1]['subida'],10.42);self.assertAlmostEqual(resumen[0]['bajada'],-3.3)

    def test_receta(self):
        p=Perfil();p.validar();self.assertEqual(p.tiempos,(248.,60.,132.))
        self.assertEqual(p.punto(5),(22.5,'Ascenso'));self.assertEqual(p.punto(248),(630,'Mantenimiento'))
        self.assertEqual(p.punto(308),(630,'Descenso'));self.assertEqual(p.punto(440),(300,'Programa finalizado'))
        self.assertEqual(p.punto(800)[0],300);p.desfase=30;self.assertEqual(p.punto(29)[1],'Antes del programa')
        p.subida=0
        with self.assertRaises(ValueError):p.validar()

    def test_asincronia_ausencias_huecos(self):
        t=datetime(2026,1,1)
        r=Registro(canales=['A','B'],muestras=[(t,{'A':620}),(t+timedelta(seconds=1),{'B':630}),(t+timedelta(minutes=5),{'A':630}),(t+timedelta(minutes=5,seconds=1),{'B':635}),(t+timedelta(minutes=10),{'A':None}),(t+timedelta(minutes=15),{'A':640}),(t+timedelta(minutes=25),{'A':645})])
        r.recalcular();self.assertEqual(r.velocidades[2]['A'],2);self.assertEqual(r.velocidades[3]['B'],1)
        self.assertIsNone(r.velocidades[5]['A']);self.assertIsNone(r.velocidades[6]['A'])
        summary,alerts=Sesion(registro=r).analisis();self.assertEqual(summary[0]['racha'],5);self.assertEqual(summary[1]['racha'],5)
        self.assertEqual(sum(a[2]=='Hueco de datos' for a in alerts),1);self.assertIsNone(summary[0]['bajada'])

    def test_lotes_repetidos_y_desordenados(self):
        t=datetime(2026,1,1);r=Registro()
        self.assertEqual(r.agregar([(t+timedelta(minutes=5),'A',25),(t,'A',20),(t,'B',21)]),3);self.assertEqual(r.velocidades[1]['A'],1)
        self.assertEqual(r.agregar([(t,'A',20),(t,'B',21)]),0);self.assertEqual(r.agregar([(t,'A',19)]),1)
        self.assertEqual(len(r.muestras),2);self.assertEqual(r.velocidades[1]['A'],1.2);before=list(r.muestras)
        with self.assertRaises(ValueError):r.agregar([(t,'A',float('nan'))])
        self.assertEqual(before,r.muestras)

    def test_guardar_exportar(self):
        s=Sesion(registro=cargar_xlsx(ROOT/'ejemplos'/'R59022 MultiChannel.xlsx'));s.alias[s.registro.canales[0]]='Puerta'
        s.controladores=[dict(fecha='2024-12-19 08:00:00',ideal=30,operador='Prueba',zonas=[10,None,20,None,None,None],temperaturas={s.registro.canales[0]:45},nota='=sin ejecutar')]
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'test.horno.json';s.guardar(path);copy=Sesion.abrir(path);self.assertEqual(copy.registro.muestras,s.registro.muestras);self.assertEqual(copy.controladores,s.controladores)
            exportar_csv(copy,Path(d)/'test.csv');exportar_excel(copy,Path(d)/'test.xlsx');w=load_workbook(Path(d)/'test.xlsx',data_only=False)
            self.assertEqual(w['Datos'].max_row,94);self.assertEqual(w['Estadísticas']['E2'].value,652.5)
            self.assertEqual(w['Controladores']['L2'].data_type,'s');self.assertEqual(w['Controladores']['L2'].value,'=sin ejecutar')
            self.assertEqual(len(w['Curvas']._charts),1);self.assertFalse(w._external_links);w.close()

    def test_xml_mapeado(self):
        xml='<Envelope xmlns="urn:x"><Lectura id="2"><Cuando>2026-01-01T10:00:00</Cuando><T>212</T></Lectura><Lectura id="3"><Cuando>2026-01-01T10:00:00</Cuando><T>INVALIDO</T></Lectura></Envelope>'
        mapping=dict(record='Lectura',time='Cuando',value='T',channel='@id',unit='°F');rows,errors=decodificar(xml,mapping)
        self.assertEqual(rows,[(datetime(2026,1,1,10),'2',100)]);self.assertEqual(len(errors),1);mapping['record']='Nada'
        with self.assertRaises(ValueError):decodificar(xml,mapping)
        self.assertEqual(decodificar('<Envelope><Lecturas /></Envelope>',mapping,allow_empty=True),([],[]))


if __name__=='__main__':unittest.main()
