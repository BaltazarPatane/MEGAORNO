"""Ejercita la GUI real, cinco vistas, exportaciones y cierre sin equipo físico."""
import argparse
import csv
import os
from pathlib import Path
import sys
import time
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from PySide6.QtWidgets import QApplication
from gui import App
from openpyxl import load_workbook
from models import Sesion

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,required=True);args=parser.parse_args()
    root=args.output.resolve();root.mkdir(parents=True,exist_ok=True)
    app=QApplication.instance() or QApplication([]);window=App(state_dir=root/'state');window.show();app.processEvents()
    assert window.stack.currentIndex()==0 and not window.active
    window.grab().save(str(root/'01-inicio.png'));window.open_example();app.processEvents()
    assert len(window.session.registro.muestras)==93 and window.tabs.count()==5
    for i,name in enumerate(['curvas','perfil','estadisticas','datos','avisos']):
        window.tabs.setCurrentIndex(i);app.processEvents();window.grab().save(str(root/f'{i+2:02d}-{name}.png'))
    window.tabs.setCurrentIndex(0);window.curves.mode.setCurrentText('Ambas');window.curves.ideal.setChecked(True);window.curves.band.setChecked(True);app.processEvents()
    window.grab().save(str(root/'07-ambas.png'))
    for kind,ext in [('PDF','pdf'),('EXCEL','xlsx'),('CSV','csv')]:
        path=root/('tratamiento.'+ext);window.export_to(path,kind);deadline=time.monotonic()+45
        while window.export_process is not None and time.monotonic()<deadline:app.processEvents();time.sleep(.02)
        assert window.last_export==path,window.notice.text()
        assert path.stat().st_size>100
    with open(root/'tratamiento.csv',encoding='utf-8-sig',newline='') as f:
        rows=list(csv.reader(f,delimiter=';'));assert len(rows)==94
    wb=load_workbook(root/'tratamiento.xlsx',read_only=True,data_only=True)
    assert wb['Datos'].max_row==94;assert 'Controladores' in wb.sheetnames;wb.close()
    if (root/'tratamiento.pdf').read_bytes()[:5]!=b'%PDF-':raise AssertionError('PDF inválido')
    original=list(window.session.registro.muestras);window.toggle_replay();window.replay_timer.stop()
    while window.replay_rows:window.replay_step()
    window.replay_step();assert window.session.registro.muestras==original
    stored=window.session_path;window.logout();assert not window.active and len(Sesion.abrir(stored).registro.muestras)==93
    window.close();print('GUI OK: login, 5 solapas, 93 muestras, reproducción, PDF, Excel, CSV, persistencia y cierre.')

if __name__=='__main__':main()
