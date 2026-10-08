"""Lectura, receta, estadística y persistencia del visor de horno."""
from dataclasses import dataclass, field, asdict
from datetime import datetime, date, time
from pathlib import Path
import csv
import json
import math
import os
import re
import tempfile
from openpyxl import load_workbook


def numero(value):
    n=float(str(value).strip().replace(',','.'))
    if not math.isfinite(n):raise ValueError('Se necesita un número finito.')
    return n


def fecha_local(value):
    if isinstance(value,datetime):dt=value
    else:
        text=str(value).strip()
        try:dt=datetime.fromisoformat(text.replace('Z','+00:00'))
        except ValueError:
            dt=None
            for fmt in ('%d/%m/%Y %H:%M:%S','%d/%m/%Y %H:%M'):
                try:dt=datetime.strptime(text,fmt);break
                except ValueError:pass
            if dt is None:raise ValueError(f'Fecha inválida: {text!r}. Usá fecha y hora ISO 8601.')
    return dt.astimezone().replace(tzinfo=None) if dt.tzinfo else dt


@dataclass
class Perfil:
    inicio:float=10.
    objetivo:float=630.
    final:float=300.
    subida:float=150.
    bajada:float=150.
    mantenimiento:float=60.
    inferior:float=620.
    superior:float=640.
    desfase:float=0.
    hueco_max:float=7.5

    def validar(self):
        if not all(math.isfinite(x) for x in asdict(self).values()):raise ValueError('Los parámetros deben ser finitos.')
        if min(self.subida,self.bajada,self.hueco_max)<=0:raise ValueError('Rampas y límite de huecos deben ser mayores que cero.')
        if self.mantenimiento<0 or self.objetivo<max(self.inicio,self.final):raise ValueError('Revisá las temperaturas y la duración.')
        if not self.inferior<=self.objetivo<=self.superior:raise ValueError('El objetivo debe estar dentro de la banda.')

    @property
    def tiempos(self):
        return ((self.objetivo-self.inicio)*60/self.subida,self.mantenimiento,(self.objetivo-self.final)*60/self.bajada)

    def punto(self,minutos):
        x=minutos-self.desfase;up,hold,down=self.tiempos
        if x<0:return self.inicio,'Antes del programa'
        if x<up:return self.inicio+x*self.subida/60,'Ascenso'
        if x<up+hold:return self.objetivo,'Mantenimiento'
        if x<up+hold+down:return self.objetivo-(x-up-hold)*self.bajada/60,'Descenso'
        return self.final,'Programa finalizado'


@dataclass
class Registro:
    origen:str=''
    serie:str=''
    canales:list=field(default_factory=list)
    muestras:list=field(default_factory=list)
    gradientes:list=field(default_factory=list)
    notas:list=field(default_factory=list)
    cortes:list=field(default_factory=list)

    def hay_corte(self,antes,despues):
        """Una interrupción real separa muestras sin modificar sus valores."""
        from bisect import bisect_right
        index=bisect_right(self.cortes,antes)
        return index<len(self.cortes) and self.cortes[index]<=despues

    def recalcular(self,hueco_max=7.5):
        self.gradientes=[];previous={};cortes=set(self.cortes)
        for dt,vals in self.muestras:
            if dt in cortes:previous.clear()
            row={c:None for c in self.canales}
            for c,value in vals.items():
                # Ausente: otro canal midió a otra hora; None: muestra inválida explícita.
                if value is None:previous.pop(c,None);continue
                if c in previous:
                    prev_dt,prev_value=previous[c];delta=(dt-prev_dt).total_seconds()/60
                    if 0<delta<=hueco_max:row[c]=(value-prev_value)/delta
                previous[c]=(dt,value)
            self.gradientes.append(row)

    def agregar(self,lecturas,hueco_max=7.5):
        lote=[(fecha_local(dt),str(c).strip(),None if v is None else numero(v)) for dt,c,v in lecturas]
        if any(not c for _,c,_ in lote):raise ValueError('Lectura sin canal.')
        filas={dt:dict(vals) for dt,vals in self.muestras};cambios=0
        for dt,c,v in lote:
            if c not in self.canales:self.canales.append(c)
            row=filas.setdefault(dt,{})
            if c not in row or row[c]!=v:row[c]=v;cambios+=1
        self.muestras=sorted(filas.items())
        if cambios:self.recalcular(hueco_max)
        return cambios


def _fecha_excel(value, hour):
    """Combine date and time columns without dropping a textual date's time."""
    if isinstance(value, str):
        text=value.strip()
        try:
            value=fecha_local(text)
        except ValueError:
            try:value=datetime.strptime(text,'%d/%m/%Y')
            except ValueError:raise ValueError(f'Fecha de Excel inválida: {text!r}.') from None
    if isinstance(value,datetime) and value.time()!=time.min:
        return fecha_local(value)
    if not isinstance(value,date):
        raise ValueError('La fecha de Excel debe contener una fecha válida.')
    day=value.date() if isinstance(value,datetime) else value
    if isinstance(hour,datetime):hour=hour.timetz()
    if isinstance(hour,str):
        text=hour.strip()
        try:hour=time.fromisoformat(text) if text else time.min
        except ValueError:raise ValueError(f'Hora de Excel inválida: {text!r}.') from None
    if hour is not None and not isinstance(hour,time):
        raise ValueError('La hora de Excel debe contener una hora válida.')
    return fecha_local(datetime.combine(day,hour or time.min))


def cargar_xlsx(ruta):
    libro=load_workbook(ruta,read_only=True,data_only=True)
    try:
        candidatas=[]
        for s in libro:
            primeras=list(s.iter_rows(max_row=min(30,s.max_row),values_only=True))
            for i,row in enumerate(primeras):
                if row and str(row[0]).strip().lower() in ('fecha','date') and any(x and '°c' in str(x).lower() for x in row[2:]):
                    candidatas.append((s,i+1,primeras));break
        if len(candidatas)!=1:raise ValueError(f'Se encontraron {len(candidatas)} hojas MadgeTech en °C. Exportá un registro individual.')
        s,header,primeras=candidatas[0];prefijos=primeras[header-2] if header>1 else ();columnas=[]
        for i,title in enumerate(primeras[header-1][2:],2):
            if title and '°c' in str(title).lower():
                nombre=re.sub(r'\s*\(°C\)\s*','',str(title),flags=re.I).strip()
                pref=str(prefijos[i]) if i<len(prefijos) and prefijos[i] else f'Columna {i+1}'
                label=pref+' · '+nombre
                if label in [x[1] for x in columnas]:label+=f' ({i+1})'
                columnas.append((i,label))
        serie=next((str(row[2] or '') for row in primeras[:header-1] if len(row)>2 and ('serie' in str(row[0]).lower() or 'serial' in str(row[0]).lower())),'')
        r=Registro(Path(ruta).name,serie,[c for _,c in columnas])
        for n,row in enumerate(s.iter_rows(min_row=header+1,values_only=True),header+1):
            if not row or row[0] is None:continue
            try:dt=_fecha_excel(row[0],row[1] if len(row)>1 else None)
            except (ValueError,TypeError) as error:raise ValueError(f'Fila {n}: {error}') from error
            if r.muestras and dt<=r.muestras[-1][0]:raise ValueError(f'Fecha repetida o fuera de orden en fila {n}.')
            vals={}
            for j,c in columnas:
                v=row[j] if j<len(row) else None
                try:vals[c]=numero(v) if v is not None and v!='' else None
                except ValueError:vals[c]=None;r.notas.append(f'Fila {n}, {c}: {v!r}; marcado sin datos.')
            r.muestras.append((dt,vals))
        if not r.muestras:raise ValueError('El archivo no contiene muestras.')
        r.recalcular();return r
    finally:libro.close()


def perfil_desde_excel(ruta):
    w=load_workbook(ruta,read_only=True,data_only=True)
    try:
        s=next((s for s in w if str(s['A1'].value).lower()=='temperatura de inicio'),None)
        if s is None:raise ValueError('No se encuentra la receta de los operadores.')
        campos=dict(inicio='H1',subida='H2',objetivo='H3',bajada='H4',final='H5',mantenimiento='H6',inferior='D8',superior='D9')
        p=Perfil(**{k:numero(s[c].value) for k,c in campos.items()});p.validar();return p
    finally:w.close()


def nombre_canal(canal):
    """Names for the documented alternating ambient/thermocouple WCF positions.

    Keep numbers from imported headers (TC3 remains TC3, never renumbered TC2).
    Unknown labels and custom physical-location aliases remain unchanged.
    """
    text=str(canal).strip()
    if text.casefold()=='madge':return 'Temperatura Ambiente'
    match=re.fullmatch(r'.+ / Posición WCF (\d+)',text)
    if match:
        position=int(match[1])
        if 1<=position<=24:
            group='Temperatura Ambiente' if position%2 else 'Termocuplas'
            return f'{group} · {(position+1)//2}'
    label=text.rsplit('·',1)[-1].strip()
    match=re.fullmatch(r'(?:Madge|Ambiente|Temperatura Ambiente)\s*(\d*)',label,re.I)
    if match:return 'Temperatura Ambiente'+(f' · {int(match[1])}' if match[1] else '')
    match=re.fullmatch(r'(?:Termopar|Termocupla|Termocuplas|TC)\s*(\d+)',label,re.I)
    if match:return f'Termocuplas · {int(match[1])}'
    return canal


@dataclass
class Sesion:
    registro:Registro=field(default_factory=Registro)
    perfil:Perfil=field(default_factory=Perfil)
    alias:dict=field(default_factory=dict)
    controladores:list=field(default_factory=list)
    titulo:str='Tratamiento térmico'
    modo:str='Histórico'

    def nombre(self,canal):
        # Presentation aliases never replace the original measurement identity.
        name=self.alias.get(canal,canal)
        if name != canal:
            return 'Temperatura Ambiente' if name.strip().casefold()=='madge' else name
        return nombre_canal(canal)

    def analisis(self):
        r,p=self.registro,self.perfil;r.recalcular(p.hueco_max);resumen,alertas=[],[]
        if not r.muestras:return resumen,alertas
        t0=r.muestras[0][0];cortes=set(r.cortes)
        for c in r.canales:
            valores=[(i,v[c]) for i,(_,v) in enumerate(r.muestras) if v.get(c) is not None]
            gradientes_validos=[(i,v[c]) for i,v in enumerate(r.gradientes) if v.get(c) is not None]
            if not valores:resumen.append(dict(canal=c,n=0));continue
            imin,vmin=min(valores,key=lambda x:x[1]);imax,vmax=max(valores,key=lambda x:x[1])
            ups=[t for t in gradientes_validos if t[1]>0];downs=[t for t in gradientes_validos if t[1]<0]
            up=max(ups,key=lambda x:x[1]) if ups else (None,None)
            down=min(downs,key=lambda x:x[1]) if downs else (None,None)
            racha=record=0.;previous=None
            for i,(dt,v) in enumerate(r.muestras):
                if dt in cortes:previous=None;racha=0.
                if c not in v:continue
                temp=v[c];anterior=previous[1] if previous else None
                delta=(dt-previous[0]).total_seconds()/60 if previous else 0
                if delta>p.hueco_max:alertas.append((dt,c,'Hueco de datos',f'{delta:.2f} min; gradiente omitida'))
                if previous and 0<delta<=p.hueco_max and temp is not None and anterior is not None and p.inferior<=temp<=p.superior and p.inferior<=anterior<=p.superior:
                    racha+=delta;record=max(record,racha)
                else:racha=0.
                previous=(dt,temp) if temp is not None else None
                elapsed=(dt-t0).total_seconds()/60;_,fase=p.punto(elapsed);gradiente=r.gradientes[i].get(c)
                if gradiente is not None and elapsed>=p.desfase and (gradiente*60>p.subida or gradiente*60 < -p.bajada):alertas.append((dt,c,'Rampa fuera de límite',f'{gradiente*60:+.2f} °C/h'))
                if temp is not None and fase=='Mantenimiento' and not p.inferior<=temp<=p.superior:alertas.append((dt,c,'Fuera de banda prevista',f'{temp:.1f} °C'))
            last_i,last_v=valores[-1]
            resumen.append(dict(canal=c,n=len(valores),actual=last_v,ultima=last_v,tultima=r.muestras[last_i][0],
                minimo=vmin,tmin=r.muestras[imin][0],maximo=vmax,tmax=r.muestras[imax][0],gradiente=r.gradientes[last_i].get(c),
                subida=up[1],bajada=down[1],tsubida=r.muestras[up[0]][0] if up[0] is not None else None,
                tbajada=r.muestras[down[0]][0] if down[0] is not None else None,racha=record))
        return resumen,sorted(alertas,key=lambda x:x[0])

    def guardar(self,ruta):
        r=self.registro
        d=dict(version=3 if r.cortes else 2,titulo=self.titulo,modo=self.modo,perfil=asdict(self.perfil),alias=self.alias,controladores=self.controladores,
               registro=dict(origen=r.origen,serie=r.serie,canales=r.canales,notas=r.notas,cortes=[dt.isoformat() for dt in r.cortes],muestras=[[dt.isoformat(),v] for dt,v in r.muestras]))
        # Serialize first, then replace from the same directory: failed writes leave
        # the previous capture intact, including when a disk fills during reception.
        payload=json.dumps(d,ensure_ascii=False,indent=2,allow_nan=False)
        dest=Path(ruta);tmp=None
        try:
            with tempfile.NamedTemporaryFile(mode='w',encoding='utf-8',dir=dest.parent,
                    prefix=dest.name+'.',suffix='.tmp',delete=False) as stream:
                tmp=Path(stream.name);stream.write(payload);stream.flush();os.fsync(stream.fileno())
            os.replace(tmp,dest)
        finally:
            if tmp is not None:
                try:tmp.unlink(missing_ok=True)
                except OSError:pass

    @classmethod
    def abrir(cls,ruta):
        d=json.loads(Path(ruta).read_text(encoding='utf-8'))
        if d.get('version') not in (2,3):raise ValueError('Formato de sesión no compatible.')
        p=Perfil(**d['perfil']);p.validar();rd=d['registro']
        r=Registro(rd.get('origen',''),rd.get('serie',''),rd['canales'],[(fecha_local(dt),v) for dt,v in rd['muestras']],notas=rd.get('notas',[]))
        r.cortes=sorted(set(fecha_local(dt) for dt in rd.get('cortes',[])))
        if not set(r.cortes).issubset({dt for dt,_ in r.muestras}):raise ValueError('Interrupción sin muestra de reanudación.')
        for i,(dt,vals) in enumerate(r.muestras):
            if i and dt<=r.muestras[i-1][0]:raise ValueError('Fechas desordenadas.')
            for c,v in vals.items():
                if c not in r.canales:raise ValueError('Canal no declarado.')
                if v is not None:vals[c]=numero(v)
        r.recalcular(p.hueco_max)
        return cls(r,p,d.get('alias',{}),d.get('controladores',[]),d.get('titulo',''),d.get('modo','Histórico'))


def exportar_csv(sesion,ruta):
    r,p=sesion.registro,sesion.perfil;r.recalcular(p.hueco_max)
    with open(ruta,'w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f,delimiter=';');w.writerow(['Fecha y hora','Minutos','Ideal °C','Etapa']+[f'{sesion.nombre(c)} {u}' for c in r.canales for u in ('°C','Gradiente °C/h','Gradiente °C/min')])
        for i,(dt,vals) in enumerate(r.muestras):
            elapsed=(dt-r.muestras[0][0]).total_seconds()/60;ideal,fase=p.punto(elapsed);row=[dt.isoformat(sep=' '),round(elapsed,6),round(ideal,6),fase]
            for c in r.canales:
                v=r.gradientes[i].get(c);row += [vals.get(c),round(v*60,6) if v is not None else None,round(v,6) if v is not None else None]
            w.writerow(row)


def exportar_controladores(sesion,ruta):
    with open(ruta,'w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f,delimiter=';');w.writerow(['Fecha y hora','Programada °C','Operador']+[f'Zona {i} %' for i in range(1,7)]+[sesion.nombre(c)+' °C' for c in sesion.registro.canales]+['Incidencia'])
        for e in sesion.controladores:w.writerow([e['fecha'],e['ideal'],e['operador']]+e['zonas']+[e['temperaturas'].get(c) for c in sesion.registro.canales]+[e['nota']])
